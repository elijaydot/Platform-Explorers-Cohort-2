from pptx import Presentation
p = Presentation(r'c:/Users/Elijay/Desktop/Platform-Explorers-Cohort-2/Capstone/Platform Explorers Use Case Template.pptx')
print('slides', len(p.slides))
for i, s in enumerate(p.slides, 1):
    t = ''
    if s.shapes.title is not None and s.shapes.title.text:
        t = s.shapes.title.text.strip()
    print(f"{i}: title={t!r} shapes={len(s.shapes)}")
