from pptx import Presentation
p = Presentation(r'c:/Users/Elijay/Desktop/Platform-Explorers-Cohort-2/Capstone/Platform Explorers Use Case Template.pptx')
for i, s in enumerate(p.slides, 1):
    print(f'--- slide {i} ---')
    for j, sh in enumerate(s.shapes, 1):
        text = ''
        if hasattr(sh, 'text'):
            text = (sh.text or '').strip().replace('\n', ' | ')
        print(f'{j:02d}. name={sh.name!r} has_text={hasattr(sh, "text")} text={text!r}')
